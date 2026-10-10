"""Build the Assessment 4 report as a Word document from reports/a4/report_text.md.

The text lives in Markdown-like source so it can be read and edited without Word; this
script lays it out to the brief (Times New Roman 12 pt, 1.5 line spacing, justified),
adds the cover page, the contribution table with live word counts, the tables filled
from the results files, the figures and the references.

Source markup, one block per paragraph:
  @part Name (ID) | sections     starts a member's part (counted separately)
  # / ##                         section and subsection headings
  - item                         bullet
  !fig name | caption | width_cm figure from reports/a4/figures
  !code name | caption | width   code excerpt image (same folder)
  !table name | caption          table built here from the results files

Word counts cover body paragraphs and bullets only: headings, captions, tables, code,
the cover, the contents and the references are excluded, so the counts are a floor.
"""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
A4 = ROOT / "reports" / "a4"
FIG = A4 / "figures"
OUT = A4 / "PRT565 A4 Report Group85.docx"
FONT = "Times New Roman"
INK = RGBColor(0x1E, 0x24, 0x27)
ACCENT = RGBColor(0xA8, 0x3A, 0x0C)
MUTED = RGBColor(0x5F, 0x68, 0x6D)

MEMBERS = [("Harsh Rastogi", "386401", "Project lead"),
           ("Saira Zafar", "407193", ""),
           ("Tharushi Wimalachandra", "387594", "")]

REFERENCES = [
    "Bergstra, J., & Bengio, Y. (2012). Random search for hyper-parameter optimization. *Journal of Machine Learning Research, 13*, 281–305.",
    "Breiman, L. (2001). Random forests. *Machine Learning, 45*(1), 5–32. https://doi.org/10.1023/A:1010933404324",
    "Breiman, L., Friedman, J. H., Olshen, R. A., & Stone, C. J. (1984). *Classification and regression trees*. Wadsworth.",
    "Chollet, F., & others. (2015). *Keras* [Computer software]. https://keras.io",
    "Davies, D. L., & Bouldin, D. W. (1979). A cluster separation measure. *IEEE Transactions on Pattern Analysis and Machine Intelligence, PAMI-1*(2), 224–227. https://doi.org/10.1109/TPAMI.1979.4766909",
    "Didan, K. (2021). *MODIS/Terra vegetation indices 16-day L3 global 250 m SIN grid V061* [Data set]. NASA EOSDIS Land Processes Distributed Active Archive Center. https://doi.org/10.5067/MODIS/MOD13Q1.061",
    "Friedman, J. H. (2001). Greedy function approximation: A gradient boosting machine. *The Annals of Statistics, 29*(5), 1189–1232. https://doi.org/10.1214/aos/1013203451",
    "Gorelick, N., Hancher, M., Dixon, M., Ilyushchenko, S., Thau, D., & Moore, R. (2017). Google Earth Engine: Planetary-scale geospatial analysis for everyone. *Remote Sensing of Environment, 202*, 18–27. https://doi.org/10.1016/j.rse.2017.06.031",
    "He, H., & Garcia, E. A. (2009). Learning from imbalanced data. *IEEE Transactions on Knowledge and Data Engineering, 21*(9), 1263–1284. https://doi.org/10.1109/TKDE.2008.239",
    "Hochreiter, S., & Schmidhuber, J. (1997). Long short-term memory. *Neural Computation, 9*(8), 1735–1780. https://doi.org/10.1162/neco.1997.9.8.1735",
    "Jeffrey, S. J., Carter, J. O., Moodie, K. B., & Beswick, A. R. (2001). Using spatial interpolation to construct a comprehensive archive of Australian climate data. *Environmental Modelling & Software, 16*(4), 309–330. https://doi.org/10.1016/S1364-8152(01)00008-1",
    "Karim, A. (2026). *PRT565 Machine learning, artificial intelligence and algorithms: Lecture slides and tutorial notebooks, Weeks 1–11* [Unpublished course materials]. Charles Darwin University.",
    "Ke, G., Meng, Q., Finley, T., Wang, T., Chen, W., Ma, W., Ye, Q., & Liu, T.-Y. (2017). LightGBM: A highly efficient gradient boosting decision tree. In *Advances in Neural Information Processing Systems 30* (pp. 3146–3154). Curran Associates.",
    "Kingma, D. P., & Ba, J. (2015). Adam: A method for stochastic optimization. In *Proceedings of the 3rd International Conference on Learning Representations*. https://arxiv.org/abs/1412.6980",
    "Lloyd, S. (1982). Least squares quantization in PCM. *IEEE Transactions on Information Theory, 28*(2), 129–137. https://doi.org/10.1109/TIT.1982.1056489",
    "North Australian Fire Information. (2026). *Northern Territory fire scar mapping* [Data set]. Darwin Centre for Bushfire Research, Charles Darwin University. https://firenorth.org.au",
    "Pedregosa, F., Varoquaux, G., Gramfort, A., Michel, V., Thirion, B., Grisel, O., Blondel, M., Prettenhofer, P., Weiss, R., Dubourg, V., Vanderplas, J., Passos, A., Cournapeau, D., Brucher, M., Perrot, M., & Duchesnay, É. (2011). Scikit-learn: Machine learning in Python. *Journal of Machine Learning Research, 12*, 2825–2830.",
    "Queensland Government. (2026). *SILO climate data: Data Drill* [Data set]. Long Paddock. https://www.longpaddock.qld.gov.au/silo/",
    "Quinlan, J. R. (1986). Induction of decision trees. *Machine Learning, 1*(1), 81–106. https://doi.org/10.1007/BF00116251",
    "Roberts, D. R., Bahn, V., Ciuti, S., Boyce, M. S., Elith, J., Guillera-Arroita, G., Hauenstein, S., Lahoz-Monfort, J. J., Schröder, B., Thuiller, W., Warton, D. I., Wintle, B. A., Hartig, F., & Dormann, C. F. (2017). Cross-validation strategies for data with temporal, spatial, hierarchical, or phylogenetic structure. *Ecography, 40*(8), 913–929. https://doi.org/10.1111/ecog.02881",
    "Rousseeuw, P. J. (1987). Silhouettes: A graphical aid to the interpretation and validation of cluster analysis. *Journal of Computational and Applied Mathematics, 20*, 53–65. https://doi.org/10.1016/0377-0427(87)90125-7",
    "Rumelhart, D. E., Hinton, G. E., & Williams, R. J. (1986). Learning representations by back-propagating errors. *Nature, 323*(6088), 533–536. https://doi.org/10.1038/323533a0",
    "Russell-Smith, J., Cook, G. D., Cooke, P. M., Edwards, A. C., Lendrum, M., Meyer, C. P., & Whitehead, P. J. (2013). Managing fire regimes in north Australian savannas: Applying Aboriginal approaches to contemporary global problems. *Frontiers in Ecology and the Environment, 11*(s1), e55–e63. https://doi.org/10.1890/120251",
    "Russell-Smith, J., Yates, C. P., Whitehead, P. J., Smith, R., Craig, R., Allan, G. E., Thackway, R., Frakes, I., Cridland, S., Meyer, M. C. P., & Gill, A. M. (2007). Bushfires ‘down under’: Patterns and implications of contemporary Australian landscape burning. *International Journal of Wildland Fire, 16*(4), 361–377. https://doi.org/10.1071/WF07018",
    "Saito, T., & Rehmsmeier, M. (2015). The precision-recall plot is more informative than the ROC plot when evaluating binary classifiers on imbalanced datasets. *PLOS ONE, 10*(3), Article e0118432. https://doi.org/10.1371/journal.pone.0118432",
    "Sokolova, M., & Lapalme, G. (2009). A systematic analysis of performance measures for classification tasks. *Information Processing & Management, 45*(4), 427–437. https://doi.org/10.1016/j.ipm.2009.03.002",
    "Srivastava, N., Hinton, G., Krizhevsky, A., Sutskever, I., & Salakhutdinov, R. (2014). Dropout: A simple way to prevent neural networks from overfitting. *Journal of Machine Learning Research, 15*, 1929–1958.",
    "Strobl, C., Boulesteix, A.-L., Zeileis, A., & Hothorn, T. (2007). Bias in random forest variable importance measures: Illustrations, sources and a solution. *BMC Bioinformatics, 8*, Article 25. https://doi.org/10.1186/1471-2105-8-25",
]


# ----------------------------------------------------------------------- helpers

def set_cell_shading(cell, hex_fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def set_cell_margins(table, top=60, bottom=60, left=90, right=90) -> None:
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for k, v in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        e = OxmlElement(f"w:{k}"); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa")
        mar.append(e)
    tblPr.append(mar)


def table_borders(table, colour="B9C2C7", inside=True) -> None:
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "bottom", "insideH") + (("left", "right", "insideV") if False else ()):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4" if edge != "top" else "8")
        e.set(qn("w:space"), "0"); e.set(qn("w:color"), colour if edge != "top" else "1E2427")
        if edge == "insideH" and not inside:
            e.set(qn("w:val"), "nil")
        b.append(e)
    for edge in ("left", "right", "insideV"):
        e = OxmlElement(f"w:{edge}"); e.set(qn("w:val"), "nil"); b.append(e)
    tblPr.append(b)


def para_border_bottom(p, colour="1E2427", size=6) -> None:
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    e = OxmlElement("w:bottom")
    e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(size)); e.set(qn("w:space"), "4")
    e.set(qn("w:color"), colour)
    bdr.append(e); pPr.append(bdr)


def para_shading(p, fill="F3F1EC") -> None:
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def keep_with_next(p) -> None:
    p.paragraph_format.keep_with_next = True


def add_runs(p, text: str, size=None, italic_all=False, colour=None) -> None:
    """Very small inline markup: *italic* and **bold**."""
    parts = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**"):
            r = p.add_run(part[2:-2]); r.bold = True
        elif part.startswith("*"):
            r = p.add_run(part[1:-1]); r.italic = True
        else:
            r = p.add_run(part)
        if italic_all:
            r.italic = True
        if size:
            r.font.size = Pt(size)
        if colour is not None:
            r.font.color.rgb = colour


def field(p, instr: str) -> None:
    r = p.add_run()
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin"); r._r.append(b)
    r2 = p.add_run()
    t = OxmlElement("w:instrText"); t.set(qn("xml:space"), "preserve"); t.text = instr
    r2._r.append(t)
    r3 = p.add_run()
    s = OxmlElement("w:fldChar"); s.set(qn("w:fldCharType"), "separate"); r3._r.append(s)
    r4 = p.add_run("1")
    r5 = p.add_run()
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end"); r5._r.append(e)


def words(text: str) -> int:
    text = re.sub(r"https?://\S+|github\.com/\S+", "link", text)
    return len(re.findall(r"[A-Za-z0-9][A-Za-z0-9'’.,%–\-/×]*", text))


# ------------------------------------------------------------------------ tables

def data():
    an = json.load(open(A4 / "analysis.json"))
    rob = json.load(open(ROOT / "reports" / "robustness.json"))
    return an, rob


def build_table(doc, name: str, an, rob):
    f3 = lambda v: f"{v:.3f}"
    if name == "sources":
        head = ["Source", "Provider", "What it gives", "Resolution", "Used"]
        rows = [["NAFI fire scars", "Charles Darwin University", "Month each pixel was mapped as burnt", "250 m, yearly", "2000–2025"],
                ["SILO Data Drill", "Queensland Government", "Daily rain, max temperature, VPD, radiation; elevation", "497 points, 0.5° apart, daily", "1999 to April 2026"],
                ["MODIS MOD13Q1", "NASA, via Google Earth Engine", "Vegetation greenness (NDVI)", "250 m, every 16 days", "2000 to April 2026"]]
        widths = [3.0, 3.4, 4.6, 3.0, 2.4]
    elif name == "classes":
        head = ["Part", "Years", "Rows", "Unburnt", "Early", "Late"]
        sh = lambda k: an[k]
        rows = [["Training", "2005–2019", f"{an['rows_train']:,}"] + [f"{sh('share_train')[c]:.1f}%" for c in ("unburnt", "early", "late")],
                ["Validation", "2020–2022", f"{an['rows_val']:,}"] + [f"{sh('share_val')[c]:.1f}%" for c in ("unburnt", "early", "late")],
                ["Test", "2023–2025", f"{an['rows_test']:,}"] + [f"{sh('share_test')[c]:.1f}%" for c in ("unburnt", "early", "late")],
                ["All modelling rows", "2005–2025", f"{an['rows_model']:,}"] + [f"{an['class_share_model'][c]:.1f}%" for c in ("unburnt", "early", "late")],
                ["All labels", "2000–2025", f"{an['rows_labels']:,}"] + [f"{an['class_share_labels'][c]:.1f}%" for c in ("unburnt", "early", "late")]]
        widths = [3.8, 2.6, 2.8, 2.4, 2.2, 2.2]
    elif name == "groups":
        head = ["Group", "Count", "What it holds"]
        rows = [["Fire history", "26", "Fire frequency and late-fire frequency over 3, 5, 10, 15, 20 and 25 years; years since the last fire and last late fire, with never-burnt flags; last year's state and burnt fractions; burnt two years ago; what the eight neighbouring cells did last year"],
                ["Weather", "14", "Wet-season (November to April) rainfall, maximum temperature, vapour pressure deficit and radiation; April rainfall and VPD; days since meaningful and heavy rain; the previous wet and dry seasons; rainfall anomaly and ratio against the cell's 2000–2019 normal"],
                ["Greenness", "7", "Mean, highest and lowest wet-season NDVI; April NDVI; green-up speed; April and peak anomalies against the cell's 2005–2019 normal"],
                ["Place", "3", "Longitude, latitude and elevation"]]
        widths = [2.6, 1.6, 11.4]
    elif name == "tuning":
        head = ["Model", "Search", "Tried", "Best settings", "Val F1", "Test F1 tuned", "Test F1 hand-set"]
        order = [("gradient boosting", "Gradient boosting", "Randomised"), ("LSTM", "LSTM", "Grid"),
                 ("multilayer ANN", "Multilayer ANN", "Grid"), ("random forest", "Random forest", "Randomised"),
                 ("decision tree", "Decision tree", "Randomised"), ("naive bayes", "Naive Bayes", "Grid")]
        def pretty(d):
            names = {"learning_rate": "learning rate", "max_leaf_nodes": "leaves", "min_samples_leaf": "rows per leaf",
                     "l2_regularization": "L2", "max_iter": "trees", "n_estimators": "trees", "max_depth": "depth",
                     "max_features": "features per split", "criterion": "criterion", "var_smoothing": "variance smoothing",
                     "units": "units", "drop": "dropout", "lr": "learning rate"}
            out = []
            for k, v in d.items():
                if isinstance(v, list):
                    v = "/".join(str(x) for x in v)
                if k == "var_smoothing":
                    v = f"{v:.0e}"
                out.append(f"{names.get(k, k)} {v}")
            return ", ".join(out)
        rows = []
        for k, nice, kind in order:
            t = rob["tuning"][k]
            rows.append([nice, kind, str(t["candidates"]), pretty(t["best_params"]), f3(t["val_f1_best"]),
                         f3(t["test_f1_tuned"]), f3(t["test_f1_manual"])])
        widths = [2.9, 2.3, 1.2, 4.9, 1.4, 1.6, 1.7]
    elif name == "results":
        head = ["Model", "Accuracy", "Precision", "Recall", "F1", "F1 late", "ROC-AUC", "PR-AUC"]
        rk = {"Gradient boosting": "gradient boosting", "LSTM": "LSTM", "Multilayer ANN": "multilayer ANN",
              "Random forest": "random forest", "Decision tree": "decision tree", "Naive Bayes": "naive bayes"}
        rows = []
        for n, m in an["metrics"].items():
            r = rob["models"].get(rk.get(n, ""), {})
            rows.append([("Baseline: " + n.lower()) if n not in rk else n, f3(m["accuracy"]), f3(m["precision"]),
                         f3(m["recall"]), f3(m["f1"]), f3(m["f1_class"]["late"]),
                         f3(r["roc_auc_macro"]) if r else "–", f3(r["pr_auc_macro"]) if r else "–"])
        widths = [3.9, 1.75, 1.75, 1.55, 1.4, 1.55, 1.75, 1.6]
    else:
        raise KeyError(name)

    t = doc.add_table(rows=1 + len(rows), cols=len(head))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    table_borders(t)
    set_cell_margins(t)
    for i, h in enumerate(head):
        c = t.rows[0].cells[i]
        c.text = ""
        p = c.paragraphs[0]
        r = p.add_run(h); r.bold = True; r.font.size = Pt(9.5)
        set_cell_shading(c, "EEF0F1")
    for ri, row in enumerate(rows, start=1):
        for ci, v in enumerate(row):
            c = t.rows[ri].cells[ci]
            c.text = ""
            p = c.paragraphs[0]
            r = p.add_run(v); r.font.size = Pt(10)
            if name == "results" and ri == 1:
                r.bold = True
            if name == "results" and v.startswith("Baseline"):
                r.italic = True
    for row in t.rows:
        row.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
        for ci, w in enumerate(widths):
            row.cells[ci].width = Cm(w)
            row.cells[ci].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in row.cells[ci].paragraphs:
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.space_before = Pt(0)
                if name in ("classes", "results", "tuning") and ci >= (2 if name != "tuning" else 4):
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                elif name == "results" and ci >= 1:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    # keep each table on one page: every row but the last keeps with the next
    for row in t.rows[:-1]:
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.keep_with_next = True
    # repeat header row on page breaks, never split a row
    trPr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)
    for row in t.rows:
        trPr = row._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit"); cs.set(qn("w:val"), "true"); trPr.append(cs)
    return t


# ------------------------------------------------------------------------ layout

def styles(doc) -> None:
    # Theme fonts would let Word swap Cambria/Calibri back in through linked character
    # styles and the document defaults, so every font reference is pinned to Times.
    for rf in doc.styles.element.iter(qn("w:rFonts")):
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rf.get(qn(a)) is not None:
                del rf.attrib[qn(a)]
        for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rf.set(qn(a), FONT)
    st = doc.styles["Normal"]
    st.font.name = FONT; st.font.size = Pt(12); st.font.color.rgb = INK
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    pf = st.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.space_after = Pt(6)
    for lvl, size, before, after in ((1, 15, 18, 6), (2, 12.5, 12, 3)):
        h = doc.styles[f"Heading {lvl}"]
        h.font.name = FONT; h.font.size = Pt(size); h.font.bold = True
        h.font.italic = False
        h.font.color.rgb = INK if lvl == 2 else ACCENT
        rpr = h.element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts"); rpr.append(fonts)
        for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            fonts.set(qn(a), FONT)
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if fonts.get(qn(a)) is not None:
                del fonts.attrib[qn(a)]
        h.paragraph_format.space_before = Pt(before)
        h.paragraph_format.space_after = Pt(after)
        h.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        h.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        h.paragraph_format.keep_with_next = True
    cap = doc.styles["Caption"]
    cap.font.name = FONT; cap.font.size = Pt(10.5); cap.font.italic = True; cap.font.bold = False
    cap.font.color.rgb = MUTED
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    cap.paragraph_format.line_spacing = 1.15
    cap.paragraph_format.space_after = Pt(10)
    from docx.enum.style import WD_STYLE_TYPE
    # Word's built-in TOC styles are named "toc 1"/"toc 2" with ids TOC1/TOC2; defining them
    # under exactly those names is the only way Word keeps them when it rebuilds the contents.
    from docx.oxml import parse_xml
    styles_el = doc.styles.element
    for sid, name, size, bold, before, indent in (("TOC1", "toc 1", 23, True, 100, 0),
                                                 ("TOC2", "toc 2", 22, False, 0, 340)):
        styles_el.append(parse_xml(
            f'<w:style xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            f'w:type="paragraph" w:styleId="{sid}"><w:name w:val="{name}"/><w:basedOn w:val="Normal"/>'
            f'<w:next w:val="Normal"/><w:uiPriority w:val="39"/><w:unhideWhenUsed/>'
            f'<w:pPr><w:tabs><w:tab w:val="right" w:leader="dot" w:pos="9060"/></w:tabs>'
            f'<w:spacing w:before="{before}" w:after="0" w:line="259" w:lineRule="auto"/>'
            f'<w:ind w:left="{indent}"/><w:jc w:val="left"/></w:pPr>'
            f'<w:rPr>{"<w:b/>" if bold else ""}<w:sz w:val="{size}"/></w:rPr></w:style>'))
    fh = doc.styles.add_style("Front Heading", WD_STYLE_TYPE.PARAGRAPH)
    fh.base_style = doc.styles["Normal"]
    fh.font.size = Pt(15); fh.font.bold = True; fh.font.color.rgb = ACCENT
    fh.paragraph_format.space_before = Pt(0); fh.paragraph_format.space_after = Pt(6)
    fh.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fh.paragraph_format.keep_with_next = True
    lb = doc.styles["List Bullet"]
    lb.font.name = FONT; lb.font.size = Pt(12)
    lb.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    lb.paragraph_format.space_after = Pt(2)


def pin_theme_fonts(doc) -> None:
    """Set the theme's major and minor Latin fonts to Times New Roman.

    Word writes theme references into the runs it generates itself (the contents
    entries, for one), so pinning the styles alone still lets Cambria through."""
    for part in doc.part.package.iter_parts():
        if str(part.partname).startswith("/word/theme/"):
            xml = part.blob.decode("utf-8")
            xml = re.sub(r'(<a:(?:major|minor)Font>\s*<a:latin typeface=")[^"]*"', r'\1Times New Roman"', xml)
            part._blob = xml.encode("utf-8")


def footer(section) -> None:
    section.different_first_page_header_footer = True
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run("PRT565 Assessment 4 · Group 85 · PYRANTIS")
    r.font.size = Pt(9.5); r.font.color.rgb = MUTED
    r = p.add_run("\t\t"); r.font.size = Pt(9.5)
    r = p.add_run(""); r.font.size = Pt(9.5)
    field(p, "PAGE")
    for run in p.runs:
        run.font.size = Pt(9.5); run.font.color.rgb = MUTED
    # right-aligned tab at the text width
    pPr = p._p.get_or_add_pPr()
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab"); tab.set(qn("w:val"), "right"); tab.set(qn("w:pos"), str(int(Cm(16).twips if hasattr(Cm(16), 'twips') else 9070)))
    tabs.append(tab); pPr.append(tabs)


def cover(doc) -> None:
    def line(text, size, bold=False, colour=INK, after=4, before=0, italic=False, align=WD_ALIGN_PARAGRAPH.LEFT, spacing=1.15):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_after = Pt(after); p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.line_spacing = spacing
        r = p.add_run(text); r.font.size = Pt(size); r.bold = bold; r.italic = italic
        r.font.color.rgb = colour
        return p

    line("CHARLES DARWIN UNIVERSITY", 11, bold=True, colour=MUTED, after=28)
    p = line("PRT565 Machine Learning, Artificial Intelligence and Algorithms", 13, after=2)
    line("Assessment 4 · Group Report · 40% of the unit", 13, colour=MUTED, after=36)

    logo = ROOT / "web" / "assets" / "pyrantis-logo.png"
    if logo.exists():
        pp = doc.add_paragraph(); pp.paragraph_format.space_after = Pt(6)
        pp.add_run().add_picture(str(logo), width=Cm(3.0))
    p = line("PYRANTIS", 34, bold=True, colour=ACCENT, after=2, spacing=1.0)
    line("Classifying early and late fire across the Northern Territory,\nand forecasting the 2026 season",
         16, after=34, spacing=1.15)
    para_border_bottom(doc.paragraphs[-1], colour="B9C2C7", size=4)

    t = doc.add_table(rows=6, cols=2)
    t.autofit = False
    set_cell_margins(t, 50, 50, 0, 60)
    info = [("Group", "Group 85"),
            ("Members", "\n".join(f"{n} ({i})" for n, i, _ in MEMBERS)),
            ("Project lead", "Harsh Rastogi"),
            ("Campus", "Darwin (Danala) Campus"),
            ("Lecturer", "Dr Asif Karim"),
            ("Due", "15 October 2026, 22:00 Sydney time")]
    for row, (k, v) in zip(t.rows, info):
        row.cells[0].width = Cm(3.6); row.cells[1].width = Cm(12.4)
        for c, text, bold in ((row.cells[0], k, True), (row.cells[1], v, False)):
            c.text = ""
            p = c.paragraphs[0]
            p.paragraph_format.line_spacing = 1.15; p.paragraph_format.space_after = Pt(2)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(text); r.font.size = Pt(12); r.bold = bold
            if bold:
                r.font.color.rgb = MUTED
    line("", 6, after=28)
    line("Assessment 3 recording: https://youtu.be/c_0yCsdZX2U", 11, colour=MUTED, after=1)
    line("Live site: https://pyrantis.harshlabs.workers.dev", 11, colour=MUTED, after=1)
    line("Code and data: github.com/harshrastogii/PYRANTIS (also in the submitted zip file)", 11, colour=MUTED, after=1)


def contribution_page(doc, counts: dict, segments: dict) -> None:
    doc.add_paragraph("Who did what", style="Front Heading")
    p = doc.add_paragraph("The work was divided by mutual agreement. Harsh Rastogi, as project lead, took half the report. Saira Zafar and Tharushi Wimalachandra took a quarter each. Each member wrote their own part, and all three reviewed the others' parts before submission. The word counts cover body text only. Headings, captions, tables, code, the cover, this page and the references are not included.")
    t = doc.add_table(rows=1 + len(segments) + 1, cols=3)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_borders(t); set_cell_margins(t, 70, 70, 90, 90)
    head = ["Member", "Segments written", "Words"]
    for i, h in enumerate(head):
        c = t.rows[0].cells[i]; c.text = ""
        r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(11)
        set_cell_shading(c, "EEF0F1")
    widths = [4.6, 9.4, 2.0]
    total = 0
    for ri, (who, segs) in enumerate(segments.items(), start=1):
        n = counts[who]; total += n
        vals = [who, segs, f"{n:,}"]
        for ci, v in enumerate(vals):
            c = t.rows[ri].cells[ci]; c.text = ""
            r = c.paragraphs[0].add_run(v); r.font.size = Pt(11)
    last = t.rows[-1].cells
    for ci, v in enumerate(["Total", "", f"{total:,}"]):
        last[ci].text = ""
        r = last[ci].paragraphs[0].add_run(v); r.bold = True; r.font.size = Pt(11)
    for row in t.rows:
        for ci, w in enumerate(widths):
            row.cells[ci].width = Cm(w)
            for p in row.cells[ci].paragraphs:
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_after = Pt(0)
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if ci == 2 else WD_ALIGN_PARAGRAPH.LEFT

    doc.add_paragraph("Contents", style="Front Heading").paragraph_format.page_break_before = True
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    field(p, 'TOC \\o "1-2" \\h \\z \\u')


def build() -> None:
    an, rob = data()
    src = (A4 / "report_text.md").read_text().split("\n\n")

    # first pass: word counts per member
    counts, segments, who = {}, {}, None
    for block in src:
        b = block.strip()
        if b.startswith("@part"):
            name_id, segs = [x.strip() for x in b[5:].split("|")]
            who = name_id
            counts.setdefault(who, 0)
            segments[who] = (segments.get(who, "") + "; " if who in segments else "") + segs
            continue
        if not b or b.startswith(("#", "!")):
            continue
        counts[who] += sum(words(l.lstrip("- ")) for l in b.splitlines())
    seg_text = {
        "Harsh Rastogi (386401)": "Group information and project title; problem description and motivation; dataset description; results obtained and overall understanding, including the 2026 forecast, ethics and limitations (Sections 1–3 and 7)",
        "Saira Zafar (407193)": "Dataset analysis with charts and what each one showed; feature transformation, preprocessing, feature engineering and feature selection; k-means analysis (Section 4)",
        "Tharushi Wimalachandra (387594)": "Choice and description of the six algorithms and the baselines; description of key steps (split, class weights, early stopping, tuning); overview of model evaluation measures (Sections 5 and 6)",
    }
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(2.5)
    sec.top_margin = Cm(2.3); sec.bottom_margin = Cm(2.2)
    styles(doc)
    pin_theme_fonts(doc)
    footer(sec)

    cover(doc)
    doc.add_page_break()
    contribution_page(doc, counts, seg_text)

    first_part = True
    for block in src:
        b = block.strip()
        if not b:
            continue
        if b.startswith("@part"):
            name_id, segs = [x.strip() for x in b[5:].split("|")]
            p = doc.add_paragraph()
            if first_part:
                p.paragraph_format.page_break_before = True
                first_part = False
            p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            para_shading(p)
            keep_with_next(p)
            r = p.add_run("Written by " + name_id); r.bold = True; r.font.size = Pt(11)
            r = p.add_run("   " + segs); r.font.size = Pt(11); r.font.color.rgb = MUTED
            continue
        if b.startswith("## "):
            doc.add_heading(b[3:], level=2); continue
        if b.startswith("# "):
            doc.add_heading(b[2:], level=1); continue
        if b.startswith("!fig") or b.startswith("!code"):
            kind, rest = b.split(" ", 1)
            parts = [x.strip() for x in rest.split("|")]
            name, caption = parts[0], parts[1]
            width = float(parts[2]) if len(parts) > 2 else 15.5
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.0
            keep_with_next(p)
            p.add_run().add_picture(str(FIG / f"{name}.png"), width=Cm(width))
            c = doc.add_paragraph(style="Caption"); add_runs(c, caption)
            continue
        if b.startswith("!table"):
            name, caption = [x.strip() for x in b[6:].split("|")]
            c = doc.add_paragraph(style="Caption"); add_runs(c, caption)
            c.paragraph_format.space_before = Pt(6); c.paragraph_format.space_after = Pt(4)
            keep_with_next(c)
            build_table(doc, name, an, rob)
            sp = doc.add_paragraph(); sp.paragraph_format.space_after = Pt(4)
            sp.paragraph_format.line_spacing = 1.0
            continue
        lines = b.splitlines()
        if all(l.startswith("- ") for l in lines):
            for l in lines:
                p = doc.add_paragraph(style="List Bullet"); add_runs(p, l[2:])
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            continue
        if b.endswith(":") and False:
            pass
        p = doc.add_paragraph(); add_runs(p, " ".join(lines))

    doc.add_heading("References", level=1).paragraph_format.page_break_before = True
    for ref in REFERENCES:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent = Cm(1.27)
        p.paragraph_format.first_line_indent = Cm(-1.27)
        p.paragraph_format.space_after = Pt(4)
        add_runs(p, ref)

    doc.save(OUT)
    json.dump({"counts": counts, "total": sum(counts.values())},
              open(A4 / "word_counts.json", "w"), indent=1)
    print("wrote", OUT.relative_to(ROOT))
    for k, v in counts.items():
        print(f"  {k:<34} {v:>5} words")
    print(f"  {'total':<34} {sum(counts.values()):>5}")


if __name__ == "__main__":
    build()
